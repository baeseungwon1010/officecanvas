"""Builds sample.docx: a Word document with what the .docx tests check. Word
saves no layout, so these are the inputs to the layout made here: a title and a
heading in built-in styles, a long paragraph to wrap, bold, italic and colored
runs, numbered and bulleted lists, a right-aligned tab stop, a paragraph with a
bottom border and one with shading, a 1.5-line paragraph, Korean text, a table in
"Table Grid" (whose paragraph spacing applies inside it) with merged cells and a
shaded cell, an inline picture, a page break, and a header and a footer with a
PAGE field.

Run: python3 make_sample_docx.py <outdir>  (needs python-docx and Pillow)"""
import io
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

out = sys.argv[1] if len(sys.argv) > 1 else "."

BODY = ("Word saves no line breaks, so this paragraph is broken into lines here, measured with the "
        "fonts the renderer draws with, and set at its line spacing inside the body width of the page. ")


def shade(element, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    element.append(shd)


def field(paragraph, instr):
    run = paragraph.add_run()
    for kind, text in (("begin", None), (None, instr), ("separate", None), (None, "1"), ("end", None)):
        if kind:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), kind)
            run._r.append(fc)
        elif text == instr:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = f" {instr} "
            run._r.append(it)
        else:
            t = OxmlElement("w:t")
            t.text = text
            run._r.append(t)


doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Cm(21), Cm(29.7)
section.left_margin = section.right_margin = Cm(2.5)
section.top_margin = section.bottom_margin = Cm(2.5)

section.header.paragraphs[0].text = "OfficeCanvas sample header"
footer = section.footer.paragraphs[0]
footer.add_run("Page ")
field(footer, "PAGE")

doc.add_heading("DOCX sample", level=0)
doc.add_heading("1. Text", level=1)
doc.add_paragraph(BODY * 3)
p = doc.add_paragraph()
p.add_run("Bold").bold = True
p.add_run(", ")
p.add_run("italic").italic = True
p.add_run(" and ")
red = p.add_run("red")
red.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
p.add_run(" runs.")

for item in ("First step", "Second step", "Third step"):
    doc.add_paragraph(item, style="List Number")
for item in ("A bullet", "Another bullet"):
    doc.add_paragraph(item, style="List Bullet")

p = doc.add_paragraph("Left\tRight")
p.paragraph_format.tab_stops.add_tab_stop(Cm(16), WD_TAB_ALIGNMENT.RIGHT)

p = doc.add_paragraph("A paragraph with a line under it.")
pBdr = OxmlElement("w:pBdr")
bottom = OxmlElement("w:bottom")
for k, v in (("val", "single"), ("sz", "8"), ("space", "1"), ("color", "2E75B6")):
    bottom.set(qn(f"w:{k}"), v)
pBdr.append(bottom)
p._p.get_or_add_pPr().append(pBdr)

p = doc.add_paragraph("A shaded paragraph.")
shade(p._p.get_or_add_pPr(), "FFF2CC")

p = doc.add_paragraph("Spaced one and a half lines. " * 6)
p.paragraph_format.line_spacing = 1.5

doc.add_paragraph("한글 문단은 띄어쓰기에서 줄이 나뉘고, 글꼴은 기기의 한글 글꼴로 그린다. " * 3)

doc.add_heading("2. Table", level=1)
table = doc.add_table(rows=4, cols=3)
table.style = "Table Grid"
table.alignment = WD_TABLE_ALIGNMENT.LEFT
for r, row in enumerate(table.rows):
    for c, cell in enumerate(row.cells):
        cell.text = f"R{r + 1}C{c + 1}"
table.cell(0, 0).merge(table.cell(0, 1)).text = "Merged across"
table.cell(1, 2).merge(table.cell(3, 2)).text = "Merged down"
shade(table.cell(2, 0)._tc.get_or_add_tcPr(), "D9E2F3")

doc.add_heading("3. Picture", level=1)
img = Image.new("RGB", (200, 100), (200, 30, 30))
img.paste((30, 30, 200), (100, 0, 200, 100))
buf = io.BytesIO()
img.save(buf, "PNG")
buf.seek(0)
doc.add_picture(buf, width=Cm(6))

p = doc.add_paragraph("Before the page break.")
p.add_run().add_break(WD_BREAK.PAGE)
doc.add_paragraph("After the page break, on page two.")

doc.save(f"{out}/sample.docx")
print("sample.docx written")
