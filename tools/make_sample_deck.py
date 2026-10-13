"""Builds sample_deck.pptx: a small, ordinary-looking deck (title slide, cards,
table, picture, closing slide) whose content the tests check by value — slide
background #F5F4F0, a full-height #3A5BA0 sidebar, an 18pt label above a 48pt
headline, a picture at (1in, 2in), and a dark closing slide with white text.
Converting it to .ppt with LibreOffice gives the matching binary fixture.
Run: python3 make_sample_deck.py <outdir>"""
import io
import sys

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

out = sys.argv[1] if len(sys.argv) > 1 else "."
ACCENT = RGBColor(0x3A, 0x5B, 0xA0)
PAPER = RGBColor(0xF5, 0xF4, 0xF0)
INK = RGBColor(0x22, 0x22, 0x22)
MUTED = RGBColor(0x66, 0x66, 0x66)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
blank = prs.slide_layouts[6]


def slide(background):
    s = prs.slides.add_slide(blank)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = background
    return s


def text(s, x, y, w, h, value, size, color=INK, bold=False):
    tf = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = value
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    return tf


def rect(s, kind, x, y, w, h, fill, line=None):
    sh = s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(1)
    return sh


# 1: title slide
s = slide(PAPER)
rect(s, MSO_SHAPE.RECTANGLE, 0, 0, 0.18, 7.5, ACCENT)
text(s, 1.0, 2.2, 8, 0.5, "OfficeCanvas 데모", 18, ACCENT, bold=True)
text(s, 1.0, 2.7, 10, 1.0, "샘플 발표 자료", 48, bold=True)
text(s, 1.0, 3.9, 10, 0.5, "PPTX · PPT · HWPX rendering on Android Canvas", 14, MUTED)

# 2: three cards with colored top bars
s = slide(PAPER)
text(s, 0.8, 0.5, 11, 0.8, "오늘 다룰 것 3가지", 32, bold=True)
for i, (title, body, color) in enumerate([
    ("01  구조", "슬라이드와 문서를 그림으로", RGBColor(0xB0, 0x3A, 0x2E)),
    ("02  서식", "글꼴, 색, 글머리 기호", ACCENT),
    ("03  그림", "JPEG, PNG, EMF, WMF", RGBColor(0x2E, 0x7D, 0x4F)),
]):
    x = 0.8 + i * 4.0
    rect(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 1.8, 3.6, 4.2, RGBColor(0xFF, 0xFF, 0xFF), RGBColor(0xCC, 0xCC, 0xCC))
    rect(s, MSO_SHAPE.RECTANGLE, x, 1.8, 3.6, 0.12, color)
    text(s, x + 0.3, 2.2, 3.0, 0.6, title, 22, color, bold=True)
    text(s, x + 0.3, 2.9, 3.0, 1.5, body, 14, MUTED)

# 3: table with a filled header row
s = slide(PAPER)
text(s, 0.8, 0.5, 11, 0.8, "형식별 지원 범위", 32, bold=True)
rows = [("형식", "구조", "비고"), ("PPTX", "zip + XML", "테마·레이아웃 상속"),
        ("PPT", "OLE2 바이너리", "마스터·자동 번호"), ("HWPX", "zip + XML", "저장된 줄 위치 사용")]
table = s.shapes.add_table(len(rows), 3, Inches(0.8), Inches(1.6), Inches(11.5), Inches(3.2)).table
for r, row in enumerate(rows):
    for c, value in enumerate(row):
        cell = table.cell(r, c)
        cell.text = value
        para = cell.text_frame.paragraphs[0]
        para.font.size = Pt(16)
        cell.fill.solid()
        if r == 0:
            cell.fill.fore_color.rgb = ACCENT
            para.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            para.font.bold = True
        else:
            cell.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if r % 2 else RGBColor(0xEE, 0xEC, 0xE6)
            para.font.color.rgb = INK

# 4: one picture at (1in, 2in)
s = slide(PAPER)
text(s, 0.8, 0.5, 11, 0.8, "그림", 32, bold=True)
img = Image.new("RGB", (300, 200), (0xF5, 0xF4, 0xF0))
draw = ImageDraw.Draw(img)
draw.rectangle((20, 20, 140, 180), fill=(0x3A, 0x5B, 0xA0))
draw.ellipse((160, 40, 280, 160), fill=(0xB0, 0x3A, 0x2E))
buf = io.BytesIO()
img.save(buf, "PNG")
buf.seek(0)
s.shapes.add_picture(buf, Inches(1), Inches(2), Inches(3), Inches(2))

# 5: dark closing slide
s = slide(RGBColor(0x2E, 0x4A, 0x86))
text(s, 1.0, 2.6, 10, 0.9, "감사합니다", 40, RGBColor(0xFF, 0xFF, 0xFF), bold=True)
text(s, 1.0, 3.6, 10, 1.5, "github.com/baeseungwon1010/officecanvas", 18, RGBColor(0xDD, 0xE4, 0xF2))

prs.save(f"{out}/sample_deck.pptx")
print("ok")
