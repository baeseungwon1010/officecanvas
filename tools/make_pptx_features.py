"""Builds pptx_features.pptx: one slide per .pptx-parser feature, with values the
tests check exactly (PptxParserTest, SlideRendererTest).

  1. theme styles: a default-styled rectangle (fill from <p:style> fillRef ->
     accent1) and a rectangle filled with accent2 at lumMod 50%
  2. master background (the slide sets none) and a title placeholder with no
     position of its own (taken from the layout)
  3. connectors: an elbow connector with an end arrow, a straight one with a
     start arrow
  4. preset outlines: right arrow, chevron, five-point star, rounded rectangle
  5. text: automatic numbering (started over after a bullet), superscript, a
     no-wrap box, a box PowerPoint shrank to 50% (normAutofit fontScale), a
     slide-number field
  6. pictures: a cropped PNG, an SVG-only icon, a picture inside
     mc:AlternateContent's fallback
  7. a table in PowerPoint's default style (stored by id only), and one with no
     style whose rows are saved shorter than their text
  8. effects: a hard outer shadow, an empty effect list overriding the style's
     theme shadow, white text on a highlight, a turned and flipped group whose
     member takes the group's fill, a shape turned and extruded in 3-D
  9. a hidden slide (skipped)
Run: python3 make_pptx_features.py <outdir>, or tools/build_fixtures.sh"""
import copy
import io
import sys

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

out = sys.argv[1] if len(sys.argv) > 1 else "."
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(10), Inches(7.5)
blank = prs.slide_layouts[6]
title_only = prs.slide_layouts[5]
A = "http://schemas.openxmlformats.org/drawingml/2006/main"


def png(color, size=(200, 100), split=None):
    img = Image.new("RGB", size, color)
    if split:
        img.paste(split, (size[0] // 2, 0, size[0], size[1]))
    buf = io.BytesIO()
    img.save(buf, "PNG")
    buf.seek(0)
    return buf


# 1: theme style references and color modifiers
s = prs.slides.add_slide(blank)
s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(1), Inches(4), Inches(2))  # default style only
sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(5.5), Inches(1), Inches(4), Inches(2))
spPr = sh._element.spPr
fill = etree.SubElement(spPr, qn("a:solidFill"))
spPr.remove(fill)
spPr.insert(list(spPr).index(spPr.find(qn("a:prstGeom"))) + 1, fill)
clr = etree.SubElement(fill, qn("a:schemeClr"), val="accent2")
etree.SubElement(clr, qn("a:lumMod"), val="50000")

# 2: master background, inherited title position
master = prs.slide_master
cSld = master._element.find(qn("p:cSld"))
bg = etree.Element(qn("p:bg"))
bgPr = etree.SubElement(bg, qn("p:bgPr"))
f = etree.SubElement(bgPr, qn("a:solidFill"))
etree.SubElement(f, qn("a:srgbClr"), val="E8F0E0")
etree.SubElement(bgPr, qn("a:effectLst"))
cSld.insert(0, bg)
s = prs.slides.add_slide(title_only)
s.shapes.title.text = "Inherited title"

# 3: connectors
s = prs.slides.add_slide(blank)
elbow = s.shapes.add_connector(MSO_CONNECTOR.ELBOW, Inches(1), Inches(1), Inches(4), Inches(3))
elbow.line.color.rgb = RGBColor(0xC0, 0, 0)
elbow.line.width = Pt(3)
etree.SubElement(elbow.line._get_or_add_ln(), qn("a:tailEnd"), type="triangle")
straight = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(5), Inches(4), Inches(9), Inches(4))
straight.line.color.rgb = RGBColor(0, 0, 0xC0)
straight.line.width = Pt(3)
etree.SubElement(straight.line._get_or_add_ln(), qn("a:headEnd"), type="triangle")

# 4: preset outlines
s = prs.slides.add_slide(blank)
for i, kind in enumerate([MSO_SHAPE.RIGHT_ARROW, MSO_SHAPE.CHEVRON, MSO_SHAPE.STAR_5_POINT, MSO_SHAPE.ROUNDED_RECTANGLE]):
    sh = s.shapes.add_shape(kind, Inches(0.3 + i * 2.4), Inches(2), Inches(2.2), Inches(2.2))
    sh.fill.solid()
    sh.fill.fore_color.rgb = RGBColor(0x20, 0x80, 0x40)

# 5: text features
s = prs.slides.add_slide(blank)
tf = s.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(4), Inches(2)).text_frame
# A bulleted paragraph between numbered ones starts the numbering over.
for i, text in enumerate(["Alpha", "Beta", "Gamma", "Note", "Delta"]):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.text = text
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", "342900")
    pPr.set("indent", "-342900")
    if text == "Note":
        etree.SubElement(pPr, qn("a:buChar"), char="•")
    else:
        etree.SubElement(pPr, qn("a:buAutoNum"), type="arabicPeriod")
tf = s.shapes.add_textbox(Inches(5), Inches(0.5), Inches(4), Inches(1)).text_frame
p = tf.paragraphs[0]
r = p.add_run(); r.text = "E=mc"
r = p.add_run(); r.text = "2"; r.font._rPr.set("baseline", "30000")
box = s.shapes.add_textbox(Inches(0.5), Inches(3), Inches(1), Inches(0.5))
box.text_frame.text = "This line is much longer than its box and does not wrap"
box.text_frame.word_wrap = False
shrunk = s.shapes.add_textbox(Inches(5), Inches(3), Inches(4), Inches(1))
shrunk.text_frame.text = "Shrunk to half"
shrunk.text_frame.paragraphs[0].runs[0].font.size = Pt(40)
bodyPr = shrunk.text_frame._txBody.find(qn("a:bodyPr"))
for child in list(bodyPr):
    bodyPr.remove(child)
etree.SubElement(bodyPr, qn("a:normAutofit"), fontScale="50000", lnSpcReduction="10000")
tf = s.shapes.add_textbox(Inches(0.5), Inches(5), Inches(3), Inches(1)).text_frame
p = tf.paragraphs[0]._p
fld = etree.SubElement(p, qn("a:fld"), id="{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}", type="slidenum")
etree.SubElement(fld, qn("a:t")).text = "5"

# 6: pictures
s = prs.slides.add_slide(blank)
pic = s.shapes.add_picture(png((255, 0, 0), split=(0, 0, 255)), Inches(0.5), Inches(1), Inches(2), Inches(2))
pic.crop_left = 0.5  # keeps the blue half
icon = s.shapes.add_picture(png((0, 0, 0)), Inches(4), Inches(1), Inches(2), Inches(2))
svg = (b'<svg viewBox="0 0 96 96" xmlns="http://www.w3.org/2000/svg">'
       b'<style>.c{fill:#00A050;}</style><circle class="c" cx="48" cy="48" r="40" fill="#000000"/></svg>')
svg_part = Part(PackURI("/ppt/media/icon1.svg"), "image/svg+xml", prs.part.package, svg)
rid = s.part.relate_to(svg_part, RT.IMAGE)
blip = icon._element.find(".//" + qn("a:blip"))
del blip.attrib[qn("r:embed")]
ext = etree.SubElement(etree.SubElement(blip, qn("a:extLst")), qn("a:ext"), uri="{96DAC541-7B7A-43D3-8B79-37D633B846F1}")
svgBlip = etree.SubElement(ext, "{http://schemas.microsoft.com/office/drawing/2016/SVG/main}svgBlip")
svgBlip.set(qn("r:embed"), rid)
fallback_pic = s.shapes.add_picture(png((255, 160, 0)), Inches(7), Inches(1), Inches(2), Inches(2))
MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
alt = etree.Element("{%s}AlternateContent" % MC, nsmap={"mc": MC})
choice = etree.SubElement(alt, "{%s}Choice" % MC, Requires="a14")
etree.SubElement(alt, "{%s}Fallback" % MC).append(copy.deepcopy(fallback_pic._element))
tree = fallback_pic._element.getparent()
tree.replace(fallback_pic._element, alt)

# 7: a table in the default style
s = prs.slides.add_slide(blank)
table = s.shapes.add_table(3, 2, Inches(1), Inches(1), Inches(8), Inches(3)).table
for r in range(3):
    for c in range(2):
        table.cell(r, c).text = f"r{r}c{c}"
# A table with no style id (no style at all) whose rows are saved far shorter
# than their text: PowerPoint grows them to fit.
plain = s.shapes.add_table(2, 1, Inches(1), Inches(4.5), Inches(4), Inches(0.2)).table
tblPr = plain._tbl.tblPr
for child in list(tblPr):
    tblPr.remove(child)
for attr in list(tblPr.attrib):
    del tblPr.attrib[attr]
for row in plain.rows:
    row.height = Emu(91440)
plain.cell(0, 0).text = "one\ntwo\nthree"
plain.cell(1, 0).text = "below"
for cell in (plain.cell(0, 0), plain.cell(1, 0)):
    for p in cell.text_frame.paragraphs:
        for run in p.runs:
            run.font.size = Pt(18)
            run.font.color.rgb = RGBColor(0, 0, 0)

# 8: effects: an explicit hard shadow down and to the right, an empty effect list
# (no shadow, despite the style's effect reference), highlighted text
s = prs.slides.add_slide(blank)
for x in (1, 5.5):
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(1), Inches(3), Inches(2))
    sh.fill.solid()
    sh.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    sh.line.fill.background()
    effects = etree.SubElement(sh._element.spPr, qn("a:effectLst"))
    if x == 1:
        shadow = etree.SubElement(effects, qn("a:outerShdw"), blurRad="0", dist="152400", dir="2700000", rotWithShape="0")
        etree.SubElement(shadow, qn("a:srgbClr"), val="000000")
tf = s.shapes.add_textbox(Inches(1), Inches(4.5), Inches(8), Inches(1)).text_frame
run = tf.paragraphs[0].add_run()
run.text = "highlighted"
run.font.size = Pt(40)
run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
hl = etree.SubElement(run._r.get_or_add_rPr(), qn("a:highlight"))
etree.SubElement(hl, qn("a:srgbClr"), val="FFC000")

# A group filled red, turned 90 degrees and flipped, whose rectangle takes the
# group's fill (<a:grpFill/>); and a rectangle turned 60 degrees about the
# vertical axis in 3-D, extruded in blue.
grp = s.shapes.add_group_shape()
inner = grp.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6), Inches(4.5), Inches(1), Inches(0.5))
inner.line.fill.background()
spPr = inner._element.spPr
for child in list(spPr):
    if child.tag in (qn("a:solidFill"), qn("a:noFill")):
        spPr.remove(child)
spPr.insert(list(spPr).index(spPr.find(qn("a:prstGeom"))) + 1, etree.Element(qn("a:grpFill")))
grp.shapes.add_shape(MSO_SHAPE.OVAL, Inches(7), Inches(5), Inches(0.5), Inches(0.5)).line.fill.background()
grpSpPr = grp._element.grpSpPr
xfrm = grpSpPr.find(qn("a:xfrm"))
xfrm.set("rot", str(90 * 60000))
xfrm.set("flipH", "1")
fill = etree.SubElement(grpSpPr, qn("a:solidFill"))
etree.SubElement(fill, qn("a:srgbClr"), val="C00000")
box3d = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(8.4), Inches(4.4), Inches(1.2), Inches(1.2))
box3d.fill.solid()
box3d.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
box3d.line.fill.background()
spPr = box3d._element.spPr
scene = etree.SubElement(spPr, qn("a:scene3d"))
camera = etree.SubElement(scene, qn("a:camera"), prst="orthographicFront")
etree.SubElement(camera, qn("a:rot"), lat="0", lon=str(60 * 60000), rev="0")
etree.SubElement(scene, qn("a:lightRig"), rig="threePt", dir="t")
sp3d = etree.SubElement(spPr, qn("a:sp3d"), extrusionH=str(Inches(0.6)))
clr = etree.SubElement(sp3d, qn("a:extrusionClr"))
etree.SubElement(clr, qn("a:srgbClr"), val="0000C0")

# 9: hidden slide
s = prs.slides.add_slide(blank)
s.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text_frame.text = "Hidden"
s._element.set("show", "0")

prs.save(f"{out}/pptx_features.pptx")
print("ok")
