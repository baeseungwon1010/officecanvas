"""Builds charts.pptx: one chart per slide, with values the tests check exactly
(ChartTest). python-pptx writes the chart parts with their value caches, the
way PowerPoint does.

  1. clustered columns: two series with fixed colors and value labels, a
     value axis fixed at 0..100 with gridlines, a legend at the bottom, a title
  2. stacked horizontal bars: three series in theme accent colors
  3. a line chart with markers
  4. a pie with its second slice colored on its own, and percent labels
  5. a doughnut
Run: python3 make_chart_deck.py <outdir>, or tools/build_fixtures.sh"""
import sys

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.util import Emu, Inches

out = sys.argv[1] if len(sys.argv) > 1 else "."

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(10), Inches(5.625)
blank = prs.slide_layouts[6]
FRAME = (Inches(1), Inches(0.5), Inches(8), Inches(4.5))


def chart(kind, categories, series):
    data = CategoryChartData()
    data.categories = categories
    for name, values in series:
        data.add_series(name, values)
    return prs.slides.add_slide(blank).shapes.add_chart(kind, *FRAME, data).chart


# 1. Clustered columns.
c = chart(XL_CHART_TYPE.COLUMN_CLUSTERED, ["1분기", "2분기", "3분기"], [("판매", (20, 40, 60)), ("목표", (30, 50, 80))])
for s, rgb in zip(c.plots[0].series, ("1F9E89", "98A2B3")):
    s.format.fill.solid()
    s.format.fill.fore_color.rgb = RGBColor.from_string(rgb)
c.plots[0].has_data_labels = True
c.plots[0].data_labels.show_value = True
c.value_axis.minimum_scale, c.value_axis.maximum_scale = 0, 100
c.value_axis.has_major_gridlines = True
c.has_legend = True
c.legend.position = XL_LEGEND_POSITION.BOTTOM
c.legend.include_in_layout = False
c.has_title = True
c.chart_title.text_frame.text = "분기별 판매"

# 2. Stacked horizontal bars.
chart(XL_CHART_TYPE.BAR_STACKED, ["A", "B"], [("x", (1, 2)), ("y", (3, 4)), ("z", (5, 6))])

# 3. A line with markers.
chart(XL_CHART_TYPE.LINE_MARKERS, ["1월", "2월", "3월", "4월"], [("방문", (10, 30, 20, 40))])

# 4. A pie with percent labels; the second slice has its own color.
c = chart(XL_CHART_TYPE.PIE, ["가", "나", "다"], [("비율", (50, 30, 20))])
point = c.plots[0].series[0].points[1]
point.format.fill.solid()
point.format.fill.fore_color.rgb = RGBColor.from_string("C00000")
c.plots[0].has_data_labels = True
c.plots[0].data_labels.show_percentage = True
c.plots[0].data_labels.show_value = False

# 5. A doughnut.
chart(XL_CHART_TYPE.DOUGHNUT, ["p", "q"], [("d", (1, 3))])

prs.save(f"{out}/charts.pptx")
print("charts.pptx: 5 slides")
